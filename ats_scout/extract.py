"""Read a resume (PDF / DOCX / TXT) and collect the raw signals the analyzer needs.

Everything here runs locally. Nothing is uploaded anywhere.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Extracted:
    path: str
    kind: str                       # "pdf" | "docx" | "txt"
    text: str = ""
    pages: int = 1
    images: int = 0
    tables: int = 0
    multi_column: bool = False
    textboxes: int = 0
    header_footer_text: str = ""
    private_use_chars: int = 0      # icon-font glyphs (U+E000..U+F8FF)
    link_count: int = 0
    font_families: set = field(default_factory=set)
    notes: list = field(default_factory=list)


_PUA = re.compile("[-]")


def extract(path: str) -> Extracted:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        ex = _extract_pdf(p)
    elif ext == ".docx":
        ex = _extract_docx(p)
    elif ext in (".txt", ".md"):
        ex = Extracted(path=str(p), kind="txt", text=p.read_text(encoding="utf-8", errors="replace"))
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    ex.private_use_chars = len(_PUA.findall(ex.text))
    ex.text = _PUA.sub("", ex.text)
    return ex


# --------------------------------------------------------------------------- PDF
def _extract_pdf(p: Path) -> Extracted:
    import pdfplumber

    ex = Extracted(path=str(p), kind="pdf")
    parts: list[str] = []
    col_votes = 0
    with pdfplumber.open(str(p)) as pdf:
        ex.pages = len(pdf.pages)
        for page in pdf.pages:
            parts.append(_page_text(page))
            ex.images += len(page.images)
            try:
                ex.tables += sum(1 for t in page.find_tables() if _is_real_table(t))
            except Exception:
                pass
            try:
                ex.link_count += len(page.hyperlinks)
            except Exception:
                pass
            for ch in page.chars[:3000]:
                name = str(ch.get("fontname", "")).split("+")[-1]
                fam = re.split(r"[-,]", name)[0]
                if fam:
                    ex.font_families.add(fam)
            try:
                if _looks_multi_column(page):
                    col_votes += 1
            except Exception:
                pass
    ex.text = "\n".join(parts)
    ex.multi_column = col_votes >= max(1, (ex.pages + 1) // 2)
    return ex


def _is_real_table(t) -> bool:
    """Design shapes (coloured bars, boxes) are often detected as 'tables'. Keep only grids full of text."""
    try:
        rows = t.extract()
    except Exception:
        return True
    cells = [c for r in rows for c in r]
    if len(rows) < 2 or max((len(r) for r in rows), default=0) < 2 or not cells:
        return False
    return sum(1 for c in cells if c and str(c).strip()) >= 0.5 * len(cells)


def _page_text(page) -> str:
    """Text of one page. Two-column pages are read column by column (main column first)."""
    try:
        from .layout import column_text
        txt = column_text(page)
        if txt is not None:
            return txt
    except Exception:
        pass
    return page.extract_text(x_tolerance=2, y_tolerance=3) or ""


def _looks_multi_column(page) -> bool:
    words = page.extract_words(x_tolerance=2, y_tolerance=3)
    return find_gutter(words) is not None


def find_gutter(words: list) -> float | None:
    """x position of a vertical white channel that most text respects (two-column page), else None."""
    if len(words) < 40:
        return None
    x_min = min(w["x0"] for w in words)
    x_max = max(w["x1"] for w in words)
    width = x_max - x_min
    if width < 100:
        return None
    bins = 100
    cover = [0] * bins
    for w in words:
        a = int((w["x0"] - x_min) / width * (bins - 1))
        b = int((w["x1"] - x_min) / width * (bins - 1))
        for i in range(a, b + 1):
            cover[i] += 1
    peak = max(cover)
    thresh = max(2, int(peak * 0.04))
    best = 0
    run = 0
    best_start = 0
    for i in range(18, 82):
        if cover[i] <= thresh:
            run += 1
            if run > best:
                best, best_start = run, i - run + 1
        else:
            run = 0
    if best < 3:
        return None
    mid = x_min + (best_start + best / 2) / bins * width
    left = sum(1 for w in words if w["x1"] < mid)
    right = sum(1 for w in words if w["x0"] > mid)
    total = len(words)
    return mid if (left > 0.15 * total and right > 0.15 * total) else None


# --------------------------------------------------------------------------- DOCX
def _extract_docx(p: Path) -> Extracted:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(p))
    ex = Extracted(path=str(p), kind="docx")
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            ex.tables += 1
            tbl = Table(child, doc)
            for row in tbl.rows:
                seen = set()
                cells = []
                for c in row.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
                lines.append("  ".join(x for x in cells if x))
    ex.text = "\n".join(lines)

    xml = doc.element.xml
    ex.images = xml.count("<pic:pic") or len(doc.inline_shapes)
    ex.textboxes = xml.count("txbxContent") // 2 or xml.count("txbxContent")

    hf = []
    for s in doc.sections:
        for part in (s.header, s.footer):
            try:
                hf.extend(par.text for par in part.paragraphs if par.text.strip())
            except Exception:
                pass
    ex.header_footer_text = "\n".join(hf)

    fams = set()
    for par in doc.paragraphs:
        for r in par.runs:
            if r.font.name:
                fams.add(r.font.name)
    ex.font_families = fams
    ex.link_count = xml.count("w:hyperlink")
    ex.pages = max(1, round(len(ex.text.split()) / 450))  # DOCX has no fixed pages; estimate
    ex.notes.append("DOCX page count is an estimate (about 450 words per page).")
    return ex
