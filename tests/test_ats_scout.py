import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import pytest

from ats_scout.analyzer import analyze, has_skill_bars, heading_key, match_job
from ats_scout.extract import Extracted, extract
from ats_scout.pdf_builder import build_fixed_cv, build_report, clean_line
from ats_scout.layout import read_pdf_layout
from ats_scout.pdf_builder import build_cv
from ats_scout.tailor import TailorRequest, suggestable_keywords
from make_samples import make, make_styled

JD = """Senior Backend Developer
Requirements:
- 3+ years of Python and Django
- PostgreSQL, REST API design, Docker, Kubernetes, AWS
Nice to have:
- Kafka, GraphQL
"""


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    return make(tmp_path_factory.mktemp("samples"))


def test_clean_cv_scores_high(samples):
    good, _ = samples
    a = analyze(extract(str(good)))
    assert a.ats_score >= 85
    assert not [i for i in a.issues if i.severity == "critical"]
    assert {k for k, _ in a.detected_sections} >= {"summary", "experience", "education", "skills"}


def test_two_column_cv_is_flagged(samples):
    _, bad = samples
    a = analyze(extract(str(bad)))
    assert a.extracted.multi_column
    assert any("column" in i.title.lower() for i in a.issues)
    assert a.ats_score < 60


def test_job_match_found_and_missing(samples):
    good, _ = samples
    a = analyze(extract(str(good)), JD)
    found = {k.keyword for k in a.match.matched}
    missing = {k.keyword for k in a.match.missing}
    assert {"python", "django", "docker"} <= found
    assert {"kubernetes", "kafka"} <= missing
    assert 0 < a.match.score < 100
    assert a.combined_score is not None


def test_synonyms_and_turkish_folding():
    m = match_job("Kubernetes ve K8s deneyimi. JS ile gelistirme. Makine Öğrenmesi projeleri.",
                  "We need Kubernetes, JavaScript and machine learning skills. Kubernetes is key.")
    assert not m.missing, [k.keyword for k in m.missing]


def test_skill_bar_detection_ignores_normal_percentages():
    assert has_skill_bars("Python ★★★★☆")
    assert has_skill_bars("Excel 80%")
    assert not has_skill_bars("- Reduced load time by 40% across 3 services")


def test_heading_detection_turkish_and_english():
    assert heading_key("İŞ DENEYİMİ") == "experience"
    assert heading_key("Work Experience:") == "experience"
    assert heading_key("EĞİTİM") == "education"
    assert heading_key("My Journey") is None


def test_fixed_pdf_is_clean_and_keeps_turkish(tmp_path):
    ex = Extracted(path="x.txt", kind="txt", text=(
        "Çağrı Şahin\ncagri@example.com | +90 555 111 22 33\nÖZET\nYazılım geliştirici, 3 yıl deneyim.\n"
        "İŞ DENEYİMİ\n• Ürün geliştirdim, %30 hızlandırdım\nEĞİTİM\nSelçuk Üniversitesi, Lisans 2020\n"
        "YETENEKLER\nPython ★★★★☆\nDocker 80%"))
    a = analyze(ex)
    out = tmp_path / "fixed.pdf"
    build_fixed_cv(a, str(out))
    again = extract(str(out))
    assert "★" not in again.text and "80%" not in again.text
    assert "Python" in again.text and "Docker" in again.text
    assert "EXPERIENCE" in again.text and not again.multi_column
    build_report(a, str(tmp_path / "report.pdf"))


def test_clean_line_strips_decorations():
    assert clean_line("Python ★★★★☆") == "Python"
    assert clean_line("SQL - 70%") == "SQL"
    assert clean_line("📧 mail@x.com") == "mail@x.com"


JD_FRONT = """Frontend Engineer
Requirements:
- JavaScript, TypeScript, React, Next.js, REST API integration
- Docker, Kubernetes, CI/CD, Tailwind CSS, SEO
Nice to have: GraphQL
"""


@pytest.fixture(scope="module")
def styled(tmp_path_factory):
    return make_styled(tmp_path_factory.mktemp("styled"))


def test_percent_results_survive_cleaning():
    assert clean_line("Cut cloud costs by 40%") == "Cut cloud costs by 40%"
    assert clean_line("Excel 80%") == "Excel"


def test_layout_reader_keeps_look(styled):
    lay = read_pdf_layout(str(styled))
    first = lay.blocks[0]
    assert first.align == "center" and first.runs[0].bold and first.size > 18
    contact = lay.blocks[1]
    assert contact.align == "center" and "ckayacann@gmail.com" in contact.text      # wrapped lines joined into one paragraph
    heads = [b for b in lay.blocks if b.heading]
    assert [b.section for b in heads] == ["summary", "experience", "education", "skills"]
    assert all(b.rule_after > 0 for b in heads)
    assert sum(b.bullet for b in lay.blocks) == 4
    summary = next(b for b in lay.blocks if b.section == "summary" and not b.heading)
    assert "\n" not in summary.text and summary.text.count("  ") == 0                # no hard line breaks inside a paragraph


def test_fixed_cv_keeps_layout_and_content(styled, tmp_path):
    a = analyze(extract(str(styled)))
    out = tmp_path / "fixed.pdf"
    res = build_cv(a, str(out))
    assert res.faithful
    again = read_pdf_layout(str(out))
    assert again.blocks[0].align == "center" and again.blocks[0].runs[0].bold
    assert [b.heading for b in again.blocks].count(True) == 4
    assert extract(str(out)).pages == 1
    assert "Laravel" in extract(str(out)).text and "40%" in extract(str(out)).text


def test_tailoring_is_honest_and_raises_match(styled, tmp_path):
    a = analyze(extract(str(styled)), JD_FRONT)
    plain = tmp_path / "plain.pdf"
    tail = tmp_path / "tail.pdf"
    build_cv(a, str(plain), TailorRequest(JD_FRONT, a.match, []))
    items = dict((t, d) for d, t in suggestable_keywords(a.match, JD_FRONT))
    assert "kubernetes" in items
    res = build_cv(a, str(tail), TailorRequest(JD_FRONT, a.match, [items["kubernetes"]]))
    t_plain, t_tail = extract(str(plain)).text, extract(str(tail)).text
    assert "JavaScript" in t_plain and " JS" not in t_plain.replace("JS/", "")       # wording aligned to the job
    assert "Kubernetes" not in t_plain                                               # never added without a tick
    assert "Kubernetes" in t_tail                                                    # added once confirmed
    assert any("Kubernetes" in n for n in res.notes)
    m_before = a.match.score
    m_after = analyze(extract(str(tail)), JD_FRONT).match.score
    assert m_after > m_before


def test_docx_layout_is_kept(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    p = d.add_paragraph()
    p.add_run("Ayşe Öz").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("ayse@example.com | +90 555 000 11 22").alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph().add_run("EXPERIENCE").bold = True
    d.add_paragraph("Developer at Acme, Jan 2022 - Present")
    d.add_paragraph("Built APIs with Python", style="List Bullet")
    path = tmp_path / "cv.docx"
    d.save(str(path))
    res = build_cv(analyze(extract(str(path))), str(tmp_path / "out.pdf"))
    assert res.faithful
    lay = read_pdf_layout(str(tmp_path / "out.pdf"))
    assert lay.blocks[0].align == "center" and "Ayşe Öz" in lay.blocks[0].text
    assert any(b.bullet for b in lay.blocks)


def test_two_column_cv_is_reflowed_in_reading_order(samples, tmp_path):
    _, bad = samples
    a = analyze(extract(str(bad)))
    assert a.extracted.multi_column
    out = tmp_path / "fixed.pdf"
    res = build_cv(a, str(out))
    assert res.reflowed
    again = extract(str(out))
    assert not again.multi_column
    assert any(b.heading for b in read_pdf_layout(str(out)).blocks)
    assert "*****" not in again.text and "80%" not in again.text
    assert analyze(again).ats_score > a.ats_score


def test_project_lines_move_out_of_experience():
    from ats_scout.layout import Block, Run, Layout
    from ats_scout.pdf_builder import _move_projects
    H = lambda txt, k: Block([Run(txt, bold=True)], heading=True, section=k)
    P = lambda txt, k, bold=False: Block([Run(txt, bold=bold)], section=k)
    L = lambda n, d: Block([Run(n, bold=True), Run(d)], section="experience")
    lay = Layout([H("EXPERIENCE", "experience"), P("Built apps; some of them:", "experience"),
                  L("Saybeni", ", numerology app."), L("Nurederm", ", care products app"),
                  H("EDUCATION", "education"), P("BSc", "education")], 600, 800, (40, 40, 40, 40))
    notes = []
    _move_projects(lay, False, notes)
    keys = [(b.section, b.heading) for b in lay.blocks]
    assert ("projects", True) in keys
    exp_text = " ".join(b.text for b in lay.blocks if b.section == "experience")
    assert "Saybeni" not in exp_text and "(see Projects)" in exp_text
    proj = [b.text for b in lay.blocks if b.section == "projects" and not b.heading]
    assert len(proj) == 2 and proj[0].startswith("Saybeni")
    assert [b.section for b in lay.blocks if b.heading] == ["experience", "projects", "education"]


def test_tracked_letters_and_split_words():
    from ats_scout.layout import fix_tracking
    mk = lambda ch, x: {"text": ch, "x0": x, "x1": x + 7, "top": 10, "bottom": 24, "size": 14}
    xs = [0, 10, 20, 30, 40, 57, 67]                      # "About" + wide gap + "me"
    ws = fix_tracking([mk(c, x) for c, x in zip("Aboutme", xs)])
    assert [w["text"] for w in ws] == ["About", "me"]
